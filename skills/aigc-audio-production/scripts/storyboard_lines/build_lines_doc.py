#!/usr/bin/env python3
"""Check first, apply confirmed punctuation/number proposals, export pure dialogue.

Python 3.10+. TXT needs no dependencies; XLSX requires openpyxl.
"""
import argparse
import hashlib
import html
import json
import re
import sys
from collections import OrderedDict
from pathlib import Path

from dialogue_rules import clean_punctuation, convert_numbers, numeric_review_reason


RULE_VERSION = 'desktop-subset-2026-10-10-v1'


def find_col(headers, keywords):
    # Exact names have priority, especially 对应台词 over other 台词 columns.
    for keyword in keywords:
        for i, value in enumerate(headers):
            if str(value or '').strip() == keyword:
                return i
    for keyword in keywords:
        for i, value in enumerate(headers):
            if keyword in str(value or ''):
                return i
    return None


def split_speaker(line):
    """Strip only plausible labels; ambiguous colon-bearing text stays intact."""
    m = re.match(r'^([^:：]{1,40})[:：]\s*(.*)$', line)
    if not m:
        return '', line
    label = m[1].strip()
    if re.search(r'[\d?!,;“”"「」]', label):
        return '', line
    if re.fullmatch(r'[\u4e00-\u9fff]{1,12}(?:[（(](?:OS|VO)[）)])?', label):
        return label, m[2].strip()
    if re.fullmatch(r'[A-ZÀ-Þ][\wÀ-ÿ.\'-]*(?:\s+[A-ZÀ-Þ][\wÀ-ÿ.\'-]*){0,3}(?:\s*[（(](?:OS|VO)[）)])?', label):
        # First-person sentences are not speaker labels.
        if label not in {'I', 'We', 'You', 'They', 'He', 'She', 'It'}:
            return label, m[2].strip()
    return '', line


def read_records(path, sheet=None, episodes=None, strip_speakers=False):
    whitelist = None
    if episodes:
        m = re.fullmatch(r'(\d+)(?:-(\d+))?', episodes)
        if not m or int(m[1]) > int(m[2] or m[1]):
            raise ValueError('--episodes 应为单集或递增范围，如 3 或 1-5')
        whitelist = set(range(int(m[1]), int(m[2] or m[1]) + 1))
    result = []
    if path.suffix.lower() == '.txt':
        if sheet or episodes:
            raise ValueError('TXT 不支持工作表或集数过滤，请先按集拆分文本')
        for row, raw in enumerate(path.read_text(encoding='utf-8-sig').splitlines(), 1):
            if not raw.strip():
                continue
            speaker, text = split_speaker(raw.strip()) if strip_speakers else ('', raw.strip())
            if not text:
                raise ValueError(f'第 {row} 行只有角色标签，没有台词')
            result.append(dict(id=f'text:{row}:1', episode='台词', sheet='', row=row,
                               line=1, shot='', speaker=speaker, original=text, context='', source_line=raw))
        actual_sheet = ''
    elif path.suffix.lower() == '.xlsx':
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        try:
            ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
            actual_sheet = ws.title
            rows = ws.iter_rows(values_only=True)
            headers = next(rows, ())
            col_ch = find_col(headers, ['章节', '集数', '集'])
            col_sb = find_col(headers, ['故事板', '分镜编号', '分镜', '板号'])
            col_lines = find_col(headers, ['对应台词', '台词'])
            col_prompt = find_col(headers, ['提示词', 'prompt'])
            if col_lines is None:
                raise ValueError('未找到台词列；不能将提示词整格当成台词，请先确认提取范围')
            if whitelist is not None and col_ch is None:
                raise ValueError('没有章节列，无法应用集数过滤')
            last_ch = None
            for rownum, row in enumerate(rows, 2):
                value = row[col_ch] if col_ch is not None else None
                if value not in (None, ''):
                    last_ch = value
                if last_ch is None:
                    episode = '台词'
                    if col_ch is not None and row[col_lines]:
                        raise ValueError(f'第 {rownum} 行无法确定章节，请确认')
                else:
                    m = re.fullmatch(r'(?:第\s*)?(\d+)(?:\.0|\s*集)?', str(last_ch).strip())
                    if not m:
                        raise ValueError(f'第 {rownum} 行章节号无法识别: {last_ch}')
                    number = int(m[1])
                    if whitelist is not None and number not in whitelist:
                        continue
                    episode = f'第{number}集'
                raw = row[col_lines]
                if raw is None or str(raw).strip() == '':
                    continue
                for linenum, raw_line in enumerate(str(raw).splitlines(), 1):
                    if not raw_line.strip():
                        continue
                    speaker, text = split_speaker(raw_line.strip())
                    if not text:
                        raise ValueError(f'第 {rownum} 行第 {linenum} 条只有角色标签')
                    result.append(dict(id=f'{actual_sheet}:{rownum}:{linenum}', episode=episode,
                        sheet=actual_sheet, row=rownum, line=linenum,
                        shot=str(row[col_sb] or '') if col_sb is not None else '',
                        speaker=speaker, original=text,
                        context=str(row[col_prompt] or '') if col_prompt is not None else '', source_line=raw_line))
        finally:
            wb.close()
    else:
        raise ValueError('只接受 .xlsx 或 UTF-8 .txt 文件')
    if not result:
        raise ValueError('未找到台词，请检查列名、文件和集数范围')
    return result, actual_sheet


DIGIT_TAGS = [
    ('年份', r'\b(?:19|20)\d{2}\b'),
    ('金额', r'[$€£¥]|dollars?|reais|pesos?|dólares?|reals?'),
    ('时间', r'a\.m\.|p\.m\.|o\'clock|\b\d{1,2}\s*(?:点|时|時)'),
    ('编号', r'\b(?:case|id|code|serial|no\.?|编号)\b'),
]


def digit_tags(text):
    """按文本特征给含数字的台词打类型标签（供统计表格展示参考，代理可微调）。"""
    tags = []
    for tag, pattern in DIGIT_TAGS:
        if re.search(pattern, text, re.I) and tag not in tags:
            tags.append(tag)
    return tags or ['普通数字']


def build_summary(path, sheet, args, items):
    episode_numbers = []
    for item in items:
        m = re.fullmatch(r'第(\d+)集', item['episode'])
        if m:
            episode_numbers.append(int(m.group(1)))
    if episode_numbers:
        lo, hi = min(episode_numbers), max(episode_numbers)
        ep_label = f'第{lo}-{hi}集' if lo != hi else f'第{lo}集'
    else:
        ep_label = '台词'
    total = len(items)
    punct = sum(i['punctuation_text'] != i['original'] for i in items)
    digit_items = [i for i in items if re.search(r'\d', i['original'])]
    tags = []
    for i in digit_items:
        for t in digit_tags(i['original']):
            if t not in tags:
                tags.append(t)
    return {
        'file': path.name,
        'episodes': ep_label,
        'total': total,
        'punctuation_suggestions': punct,
        'digit_count': len(digit_items),
        'digit_tags': tags,
        'digit_label': f"{len(digit_items)}条" + (f"（{'/'.join(tags)}）" if tags else ''),
    }


def make_review(args):
    path = Path(args.input).resolve()
    records, sheet = read_records(path, args.sheet, args.episodes, args.strip_speakers)
    items = []
    for record in records:
        punctuation = clean_punctuation(record['original'])
        numbers, reason = convert_numbers(record['original'], args.language, record['context'])
        combined, _ = convert_numbers(punctuation, args.language, record['context'])
        notes = []
        if punctuation != record['original']:
            notes.append('标点/空格清理：查看原文与建议的逐处差异；大小写和用词保持原样。')
        if numbers != record['original']:
            notes.append('数字按指定语言和当前语境等值转换；请确认数量、牌面或点数用途。')
        if reason:
            notes.append(reason)
        items.append({**record, 'punctuation_text': punctuation, 'number_text': numbers,
                      'suggested_text': combined, 'number_review_reason': reason, 'notes': notes,
                      'approve_punctuation': False, 'approve_numbers': False})
    summary = build_summary(path, sheet, args, items)
    return dict(version=RULE_VERSION, source=dict(path=str(path), sha256=hashlib.sha256(path.read_bytes()).hexdigest(),
        sheet=sheet, episodes=args.episodes, language=args.language, strip_speakers=args.strip_speakers),
        summary=summary, items=items)


def review_markdown(review):
    """已废弃：检查不再生成 md 总结文档，只输出统计表格。保留空实现防止误调。"""
    return ''


def apply_review(expected, approved):
    if not isinstance(approved, dict) or set(approved) != set(expected):
        raise ValueError('确认清单结构不一致')
    if approved.get('version') != expected['version'] or approved.get('source') != expected['source']:
        raise ValueError('来源文件、规则版本或处理范围已变更，请重新检查并确认')
    if not isinstance(approved.get('items'), list) or len(approved['items']) != len(expected['items']):
        raise ValueError('确认清单不能增删台词')
    eps = OrderedDict()
    for original, item in zip(expected['items'], approved['items']):
        if not isinstance(item, dict) or set(item) != set(original):
            raise ValueError('条目结构不一致')
        for key, value in original.items():
            if key not in {'approve_punctuation', 'approve_numbers'} and item.get(key) != value:
                raise ValueError(f"条目 {original['id']} 被改写；只能改变确认开关")
        if any(type(item.get(key)) is not bool for key in ['approve_punctuation', 'approve_numbers']):
            raise ValueError('确认开关必须为 true/false')
        text = original['punctuation_text'] if item['approve_punctuation'] else original['original']
        if item['approve_numbers']:
            # Check original context before applying punctuation to avoid changing classification.
            if original['number_review_reason']:
                raise ValueError(f"{original['id']}: 数字尚需人工确认，脚本不支持直接批准转换")
            text, reason = convert_numbers(text, expected['source']['language'], original['context'])
            if reason:
                raise ValueError(reason)
        eps.setdefault(original['episode'], []).append(text)
    return eps


def build_xml(eps, title='台词汇总'):
    parts = [f'<title>{html.escape(title)}</title>', '<p>每组一个代码块，可复制完整台词。</p>', '<hr/>']
    for episode, lines in eps.items():
        parts.append(f'<h1>{html.escape(episode)}（{len(lines)} 条台词）</h1>')
        parts.append(f'<pre lang="text" caption="{html.escape(episode, quote=True)} 台词全文"><code>'
                     + html.escape('\n'.join(lines)) + '</code></pre>')
        parts.append('<hr/>')
    return '\n'.join(parts)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input')
    parser.add_argument('--mode', choices=['check', 'extract'], default='check')
    parser.add_argument('--language', choices=['en', 'pt-BR', 'es-MX', 'unknown'], default='unknown')
    parser.add_argument('--output', '-o', required=True)
    parser.add_argument('--episodes')
    parser.add_argument('--sheet')
    parser.add_argument('--strip-speakers', action='store_true', help='仅 TXT：去掉已明确的角色前缀')
    parser.add_argument('--title', default='台词汇总')
    parser.add_argument('--review', help='确认后的 JSON 清单')
    parser.add_argument('--confirmed', action='store_true', help='已取得用户确认或检查通过；不能替代真实确认')
    args = parser.parse_args()
    try:
        output = Path(args.output).resolve()
        suffix = '.json' if args.mode == 'check' else '.xml'
        if output.suffix.lower() != suffix:
            raise ValueError(f'{args.mode} 输出应使用 {suffix} 扩展名')
        companion = output.with_suffix('.md' if args.mode == 'check' else '.txt')
        inputs = {Path(args.input).resolve()}
        if args.review:
            inputs.add(Path(args.review).resolve())
        if output in inputs or companion in inputs:
            raise ValueError('输出路径不能覆盖源文件或确认清单')
        if args.mode == 'extract' and (not args.review or not args.confirmed):
            raise ValueError('提取前需提供 --review 和 --confirmed；先展示检查清单并取得确认')
        expected = make_review(args)
        if args.mode == 'check':
            output.write_text(json.dumps(expected, ensure_ascii=False, indent=2), encoding='utf-8')
            s = expected['summary']
            print(f"检查完成（未应用修改）：{s['file']} | {s['episodes']} | 台词 {s['total']}条 | "
                  f"标点建议 {s['punctuation_suggestions']}条 | 数字相关 {s['digit_label']} -> {output}",
                  file=sys.stderr)
        else:
            approved = json.loads(Path(args.review).read_text(encoding='utf-8'))
            eps = apply_review(expected, approved)
            output.write_text(build_xml(eps, args.title), encoding='utf-8')
            companion.write_text('\n\n'.join(ep + '\n' + '\n'.join(lines) for ep, lines in eps.items()) + '\n', encoding='utf-8')
            print(f'提取完成：{len(eps)} 组，{sum(map(len, eps.values()))} 条 -> {output}', file=sys.stderr)
    except (ValueError, KeyError, OSError, ImportError) as error:
        parser.exit(1, f'错误：{error}\n')


if __name__ == '__main__':
    main()
