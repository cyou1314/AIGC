#!/usr/bin/env python3
"""Delete explicitly selected single-line metadata from a plain XLSX copy.

No implicit labels, in-place writes, whitespace cleanup, network or generation.
Formula/advanced-feature workbooks are refused instead of lossy round-tripping.
"""
from pathlib import Path
import argparse
import hashlib
import json
import re
import shutil
import tempfile
import xml.etree.ElementTree as ET
import zipfile

PROTECTED = ('画面', '对应台词', '对应原文', '时长', '预估时长', '场景', '角色', '道具')
HEADERS = PROTECTED + ('人物状态与空间', '上一板结束状态锚定', '负面约束', '声音设计', '镜头衔接逻辑')


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def trim_text(text, labels):
    if not labels or any(not s or s[-1] not in ':：' for s in labels):
        raise ValueError('Supply exact nonempty labels ending in a colon')
    if any(s[:-1] in PROTECTED for s in labels):
        raise ValueError('Narrative, source, asset and timing fields are outside this metadata helper')
    lines = text.splitlines(keepends=True)
    removed, kept, in_body, body_hits = [], [], False, 0
    headings = tuple(x + c for x in HEADERS for c in ':：') + tuple(labels)
    for index, line in enumerate(lines):
        clean = line.lstrip(' \t')
        if clean.startswith(('画面：', '画面:')):
            in_body = True
        match = next((s for s in labels if clean.startswith(s)), None)
        if not match:
            kept.append(line)
            continue
        if in_body:
            body_hits += 1
            kept.append(line)
            continue
        payload = clean[len(match):]
        if any(x + c in payload for x in PROTECTED for c in ':：') or re.search(r'\d+(?:\.\d+)?秒\s*[~～—-]', payload):
            raise ValueError('A target line also contains protected content; handle it explicitly')
        following = next((s.lstrip(' \t') for s in lines[index + 1:] if s.strip()), '')
        if following and not following.startswith(headings) and not following.startswith('【'):
            raise ValueError('Possible multiline metadata continuation; no automatic deletion')
        removed.append({'line': index + 1, 'label': match, 'text': line})
    return ''.join(kept), removed, body_hits


def check_archive(source):
    with zipfile.ZipFile(source) as archive:
        entries = archive.infolist()
        if sum(x.file_size for x in entries) > 64 * 1024 * 1024 or len(entries) > 10000:
            raise ValueError('Workbook exceeds the bounded plain-table helper limits')
        forbidden = ('xl/drawings/', 'xl/media/', 'xl/charts/', 'xl/pivot', 'xl/externalLinks/',
                     'xl/ctrlProps/', 'xl/slicer', 'customXml/', 'xl/richData/', 'xl/threadedComments/')
        for entry in entries:
            if entry.filename.startswith(forbidden) or 'vbaProject' in entry.filename:
                raise ValueError('Advanced workbook objects require the native spreadsheet workflow')
            if entry.filename.startswith('xl/') and entry.filename.endswith('.xml'):
                body = archive.read(entry)
                if re.search(rb'<(?:[\w.-]+:)?extLst(?=[\s/>])', body):
                    raise ValueError('Workbook extension objects are not supported for saving')


def snapshot(book):
    sheets = []
    for sheet in book.worksheets:
        cells = {}
        for row in sheet:
            for cell in row:
                cells[cell.coordinate] = (cell.value, cell.data_type, tuple(cell._style) if cell.has_style else None,
                    cell.hyperlink.target if cell.hyperlink else None,
                    (cell.comment.text, cell.comment.author) if cell.comment else None)
        sheets.append({'title': sheet.title, 'shape': (sheet.max_row, sheet.max_column),
            'state': sheet.sheet_state, 'freeze': str(sheet.freeze_panes),
            'merged': tuple(sorted(str(x) for x in sheet.merged_cells.ranges)),
            'filter': sheet.auto_filter.ref, 'cells': cells,
            'row_dimensions': {k: (v.height, v.hidden, v.outlineLevel, v.style_id) for k, v in sheet.row_dimensions.items()},
            'column_dimensions': {k: (v.min, v.max, v.width, v.hidden, v.outlineLevel, v.style_id) for k, v in sheet.column_dimensions.items()},
            'validation': ET.tostring(sheet.data_validations.to_tree()),
            'protection': ET.tostring(sheet.protection.to_tree())})
    return sheets


def selected_rows(spec, maximum):
    if not spec:
        return list(range(2, maximum + 1))
    rows = set()
    for value in spec.split(','):
        if not re.fullmatch(r'\d+(?:-\d+)?', value.strip()):
            raise ValueError('Rows must use 2,5,7-10 format')
        ends = list(map(int, value.strip().split('-')))
        start, end = ends[0], ends[-1]
        if start < 2 or end < start or end > maximum:
            raise ValueError('Row selection is outside the actual data range')
        rows.update(range(start, end + 1))
    return sorted(rows)


def process(source, labels, *, sheet=None, column='提示词', rows=None, output=None, write=False):
    from openpyxl import load_workbook
    from openpyxl.utils import column_index_from_string
    source = Path(source).expanduser().resolve()
    if source.suffix.lower() != '.xlsx' or not source.is_file():
        raise ValueError('Provide an existing plain .xlsx file')
    labels = tuple(labels)
    trim_text('', labels)
    destination = Path(output).expanduser().resolve() if output else None
    if write and (not destination or destination == source or destination.exists() or destination.suffix.lower() != '.xlsx'):
        raise ValueError('Write requires a new .xlsx output distinct from the source')
    before_hash = sha(source)
    check_archive(source)
    book = load_workbook(source, data_only=False)
    try:
        if sum(ws.max_row * ws.max_column for ws in book) > 200000:
            raise ValueError('Workbook exceeds the bounded cell count')
        if any(cell.data_type == 'f' for ws in book for row in ws for cell in row):
            raise ValueError('Formula workbook requires a native editor to preserve formula caches')
        if sheet:
            if sheet not in book.sheetnames:
                raise ValueError('Requested sheet is unavailable')
            ws = book[sheet]
        else:
            candidates = [s for s in book if sum(str(c.value).strip() == column for c in s[1]) == 1]
            if len(book.worksheets) == 1:
                ws = book.worksheets[0]
            elif len(candidates) == 1:
                ws = candidates[0]
            else:
                raise ValueError('Multiple possible sheets; specify the actual sheet')
        if re.fullmatch('[A-Za-z]{1,3}', column):
            col = column_index_from_string(column.upper())
            if col > ws.max_column:
                raise ValueError('Column is outside the worksheet')
        else:
            cols = [c.column for c in ws[1] if str(c.value).strip() == column]
            if len(cols) != 1:
                raise ValueError('Prompt column header must be unique and exact')
            col = cols[0]
        targets = selected_rows(rows, ws.max_row)
        expected = snapshot(book)
        sheet_index = book.sheetnames.index(ws.title)
        changes, body_hits = [], 0
        for row_number in targets:
            cell = ws.cell(row_number, col)
            if cell.value in (None, ''):
                continue
            if not isinstance(cell.value, str):
                raise ValueError('Prompt cells must be plain text')
            new, removed, hits = trim_text(cell.value, labels)
            body_hits += hits
            if new != cell.value:
                changes.append({'cell': cell.coordinate, 'removed_lines': len(removed),
                                'labels': [x['label'] for x in removed]})
                cell.value = new
                old_cell = expected[sheet_index]['cells'][cell.coordinate]
                expected[sheet_index]['cells'][cell.coordinate] = (new,) + old_cell[1:]
        if snapshot(book) != expected:
            raise ValueError('Unexpected changes detected before saving')
        report = {'mode': 'written' if write else 'dry-run', 'sheet': ws.title, 'selected_rows': targets,
                  'changed': len(changes), 'changes': changes, 'body_label_occurrences_preserved': body_hits,
                  'source_sha256': before_hash, 'output': None, 'source_unchanged': sha(source) == before_hash,
                  'disk_verified': False, 'scope': 'Plain XLSX values/types/basic styles/merges/filter/freeze/dimensions/validation/protection, not rendering or director semantics'}
        if write:
            destination.parent.mkdir(parents=True, exist_ok=True)
            with tempfile.TemporaryDirectory(prefix='.aigc-trim-', dir=destination.parent) as td:
                candidate = Path(td) / 'candidate.xlsx'
                if changes:
                    book.save(candidate)
                else:
                    shutil.copy2(source, candidate)
                saved = load_workbook(candidate, data_only=False)
                try:
                    if snapshot(saved) != expected:
                        raise ValueError('Saved workbook differs outside the authorized deletion')
                finally:
                    saved.close()
                if sha(source) != before_hash:
                    raise ValueError('Source changed during processing; no result published')
                with destination.open('xb') as handle:
                    handle.write(candidate.read_bytes())
            report.update(output=str(destination), disk_verified=True, source_unchanged=sha(source) == before_hash)
        if not report['source_unchanged']:
            raise ValueError('Source changed during processing')
        report['ok'] = True
        return report
    finally:
        book.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', required=True)
    parser.add_argument('--label', action='append', required=True)
    parser.add_argument('--sheet'); parser.add_argument('--column', default='提示词')
    parser.add_argument('--rows'); parser.add_argument('--output'); parser.add_argument('--write', action='store_true')
    args = parser.parse_args()
    try:
        result = process(args.input, args.label, sheet=args.sheet, column=args.column, rows=args.rows,
                         output=args.output, write=args.write)
        print(json.dumps(result, ensure_ascii=False, indent=2)); return 0
    except (ValueError, OSError, ImportError, zipfile.BadZipFile) as exc:
        print(json.dumps({'ok': False, 'status': 'not-delivered', 'reason': str(exc),
                          'continue_now': 'Preserve the original; use a read-only change list or native spreadsheet workflow.'}, ensure_ascii=False))
        return 2


if __name__ == '__main__':
    raise SystemExit(main())
