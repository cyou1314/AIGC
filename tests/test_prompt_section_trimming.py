"""Functional protection tests for the bounded XLSX metadata operation."""
from pathlib import Path
import importlib.util
import tempfile
import unittest
from unittest.mock import patch

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
except ImportError:
    openpyxl = None

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('bounded_trim',ROOT/'skills/prompt-optimizer/scripts/trim_prompt_sections.py')
helper=importlib.util.module_from_spec(spec);spec.loader.exec_module(helper)

PROMPT='电影质感\r\n视觉锚点：@角色A\r\n\r\n\r\n站位：A在左\r\n画面：\r\n0秒~5秒：A说“站位：这里。”\r\n站位：这行属于画面正文\r\n时长：5秒\r\n'
EXPECTED='电影质感\r\n\r\n\r\n画面：\r\n0秒~5秒：A说“站位：这里。”\r\n站位：这行属于画面正文\r\n时长：5秒\r\n'


class TextRules(unittest.TestCase):
    def test_preserves_exact_nonremoved_lines_whitespace_and_body(self):
        new,deleted,hits=helper.trim_text(PROMPT,['视觉锚点：','站位：'])
        self.assertEqual(new,EXPECTED);self.assertEqual(len(deleted),2);self.assertEqual(hits,1)

    def test_rejects_multiline_continuation(self):
        with self.assertRaises(ValueError):helper.trim_text('站位：A在左\nB在右\n画面：\n动作',['站位：'])

    def test_rejects_mixed_protected_sections(self):
        with self.assertRaises(ValueError):helper.trim_text('站位：A 画面：动作\n',['站位：'])

    def test_no_default_labels_and_no_deletion_of_source_or_picture(self):
        for labels in [[],['站位'],['画面：'],['对应台词：']]:
            with self.subTest(labels=labels),self.assertRaises(ValueError):helper.trim_text(PROMPT,labels)


@unittest.skipIf(openpyxl is None,'Optional openpyxl unavailable')
class WorkbookProtection(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.directory=Path(self.tmp.name);self.source=self.directory/'source.xlsx';self.output=self.directory/'result.xlsx'
        b=openpyxl.Workbook();s=b.active;s.title='分镜';s.append(['章节','提示词','台词']);s.append([1,PROMPT,'Mãe... 21!']);s.append([1,PROMPT,'KEEP'])
        s['B2'].font=Font(bold=True,color='112233');s['B2'].fill=PatternFill('solid',fgColor='EEEEEE');s.freeze_panes='B2';s.auto_filter.ref='A1:C3';s.row_dimensions[2].height=38
        other=b.create_sheet('其他');other['A1']='untouched';other['A2']='merged';other.merge_cells('A2:B2');other['C3']='LINK';other['C3'].hyperlink='https://example.com'
        b.save(self.source);b.close();self.original=self.source.read_bytes()

    def tearDown(self):self.tmp.cleanup()

    def run_trim(self,**kwargs):
        return helper.process(self.source,['视觉锚点：','站位：'],**kwargs)

    def test_dry_run_writes_nothing(self):
        result=self.run_trim();self.assertEqual(result['changed'],2);self.assertFalse(result['disk_verified']);self.assertEqual(list(self.directory.iterdir()),[self.source]);self.assertEqual(self.source.read_bytes(),self.original)

    def test_end_to_end_new_copy_preserves_all_other_cells_and_styles(self):
        result=self.run_trim(write=True,output=self.output,rows='2')
        self.assertEqual(result['changed'],1);self.assertTrue(result['disk_verified']);self.assertEqual(self.source.read_bytes(),self.original)
        b=openpyxl.load_workbook(self.output);s=b['分镜'];self.assertEqual(s['B2'].value,EXPECTED);self.assertEqual(s['B3'].value,PROMPT);self.assertEqual(s['C2'].value,'Mãe... 21!');self.assertTrue(s['B2'].font.bold);self.assertEqual(s.freeze_panes,'B2');self.assertEqual(str(b['其他'].merged_cells),'A2:B2');self.assertEqual(b['其他']['C3'].hyperlink.target,'https://example.com');b.close()

    def test_rejects_in_place_and_existing_output(self):
        with self.assertRaises(ValueError):self.run_trim(write=True,output=self.source)
        self.output.write_bytes(b'KEEP')
        with self.assertRaises(ValueError):self.run_trim(write=True,output=self.output)
        self.assertEqual(self.output.read_bytes(),b'KEEP');self.assertEqual(self.source.read_bytes(),self.original)

    def test_rejects_formula_workbooks_without_touching_source(self):
        b=openpyxl.load_workbook(self.source);b['其他']['B1']='=1+1';b.save(self.source);b.close();original=self.source.read_bytes()
        with self.assertRaises(ValueError):self.run_trim(write=True,output=self.output)
        self.assertEqual(self.source.read_bytes(),original);self.assertFalse(self.output.exists())

    def test_ambiguous_sheet_and_invalid_row_selection_are_not_guessed(self):
        b=openpyxl.load_workbook(self.source);s=b.create_sheet('另一个分镜');s.append(['章节','提示词']);s.append([1,PROMPT]);b.save(self.source);b.close()
        with self.assertRaises(ValueError):self.run_trim()
        for rows in ['1','3-2','99']:
            with self.subTest(rows=rows),self.assertRaises(ValueError):self.run_trim(sheet='分镜',rows=rows)

    def test_saved_corruption_is_rejected_before_result_publication(self):
        save=openpyxl.workbook.workbook.Workbook.save
        def corrupt(book,path):
            book['其他']['A1']='UNEXPECTED';return save(book,path)
        with patch.object(openpyxl.workbook.workbook.Workbook,'save',corrupt),self.assertRaises(ValueError):
            self.run_trim(write=True,output=self.output)
        self.assertFalse(self.output.exists());self.assertEqual(self.source.read_bytes(),self.original)

    def test_no_hit_is_not_counted_as_a_change(self):
        result=helper.process(self.source,['不存在：'],write=True,output=self.output)
        self.assertEqual(result['changed'],0);self.assertEqual(self.output.read_bytes(),self.original)

    def test_media_objects_are_not_lossily_resaved(self):
        from openpyxl.drawing.image import Image
        from PIL import Image as PILImage
        icon=self.directory/'synthetic.png';PILImage.new('RGB',(2,2),(0,0,0)).save(icon)
        b=openpyxl.load_workbook(self.source);b['其他'].add_image(Image(str(icon)),'D1');b.save(self.source);b.close();original=self.source.read_bytes()
        with self.assertRaises(ValueError):self.run_trim(write=True,output=self.output)
        self.assertEqual(self.source.read_bytes(),original);self.assertFalse(self.output.exists())


if __name__=='__main__':unittest.main()
